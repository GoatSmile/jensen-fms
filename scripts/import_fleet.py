#!/usr/bin/env python3
"""
Fleet register -> review lists (and, later, the import migration).

Source: Dennis's register "KOMMUNE og VIRKSOMHEDS OVERSIGT(1).xlsx" — one sheet
per customer, one row per bike that is already out with that customer, plus
twelve monthly sheets that are the service-agreement RENEWAL SCHEDULE (each
bike's anniversary month, years 2/3/4 invoiced or not, and its yearly price).

    python3 scripts/import_fleet.py review            # defaults below
    python3 scripts/import_fleet.py review --src X.xlsx --out DIR
    python3 scripts/import_fleet.py plan --batch register-2026-10-pilot-allerod --sheets "ALLERØD KOMMUNE"
    python3 scripts/import_fleet.py sql  --batch … --sheets … [--by Nazar] [--migration N]

`review` writes CSVs for a human to check and touches no database. It READS
production (customers, units, existing frames and identifiers) through
`supabase db query --linked`, to propose a customer for every sheet and
department and to flag collisions before anything is loaded.

`plan` and `sql` work on ONE batch (docs/plan-register-import.md §2–§7) and
never write to a database either: `plan` writes the mapping file a person
confirms (scripts/import-batches/<batch>/mapping.json), the exceptions, the
manifest and a review for Dennis; `sql` plans again and writes the apply,
dry-run, verify and rollback files. Applying them is a separate, deliberate
step (plan §5).

Rules this encodes (docs/plan-go-live.md §6, DECISIONS 2026-09-26):
  * Sheets holding PERSONAL data — citizens' loan registers and private buyers
    (first name / last name / email columns) — are skipped whole, never read
    into the output.
  * People's names and phone numbers are left out of every output row; the
    department, site address and EAN stay (they describe the customer).
  * Column letters differ per sheet, so every column is found by its label.
  * Frame, battery and charger numbers are split over three cells in the
    register (prefix · digits · letter); they are joined into one canonical,
    space-free, uppercase string. The frame's last letter is a year code.
  * Column A is Jensen's recognition code (BKTM01) — kept only when it looks
    like one; words there ("stjålet") become a status hint instead.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import difflib
import json
import re
import subprocess
import sys
import uuid
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

DEFAULT_SRC = Path.home() / "Documents/1-Projects/Jensen/Outdated lists/KOMMUNE og VIRKSOMHEDS OVERSIGT(1).xlsx"
DEFAULT_OUT = Path.home() / f"Documents/1-Projects/Jensen/Fleet/import-review-{dt.date.today().isoformat()}"
REPO = Path(__file__).resolve().parent.parent

MONTHS = ["Januar", "Februar", "Marts", "April", "Maj", "Juni", "Juli",
          "August", "September", "Oktober", "November", "December"]
# Not bikes: contents, contacts, price/labour tables, one-off logs.
NOT_FLEET = {"INDHOLDSFORTEGNELSE", "Ark1", "Sheet1", "KONTAKTLISTE", "LAKERING",
             "ÅRSTAL", "TIMELØN", "AVANCE", "GPS ABONNOMENTER UDEN SERVICE",
             "Høje Tåstrup Batterier", *MONTHS}

LABELS = {  # normalised header text -> field
    "stel nummer": "frame", "stelnummer": "frame", "nøgle nr": "key_no", "model": "model",
    "batteri nummer": "battery", "lader nummer": "charger", "bat.nøg.nr": "battery_key_no",
    "køn": "gender", "str": "size", "nr": "seq", "udlev. dato": "delivered", "købs dato": "delivered",
    "købsdato": "delivered", "overtagelses dato": "takeover", "kunde": "department",
    "leverings sted": "site", "aftale": "agreement", "s aftale": "agreement", "gps": "gps",
    "gps tel": "gps_phone", "ean": "ean", "tekst": "text", "fakturering": "invoiced",
    # personal / contact columns: recognised so they can be DROPPED, never output
    "kontaktperson": "_contact", "kontakt person": "_contact", "telefon nr": "_phone",
    "mobil": "_phone", "telefon": "_phone", "kontakt 2": "_contact",
    "fornavn": "_personal", "efternavn": "_personal", "email": "_personal", "adresse": "_personal",
}
SPLIT3 = {"frame", "battery", "charger"}  # prefix · digits · letter across three cells

STATUS_WORDS = [  # (pattern, bike status, reason)
    (r"stjål|stjal", "lost_or_stolen", "stolen"),
    (r"kan ikke findes|bortkommet|forsvundet", "lost_or_stolen", "cannot be found"),
    (r"udgået|kasseret|skrottet", "retired", "retired"),
]
RECOGNITION = re.compile(r"^[A-ZÆØÅ]{2,}[A-ZÆØÅ0-9 ]*\d+[A-Z]?$")
DA_MONTHS = {m: i for i, m in enumerate(["januar", "februar", "marts", "april", "maj", "juni", "juli",
                                          "august", "september", "oktober", "november", "december"], 1)}


def norm(s) -> str:
    return re.sub(r"\s+", " ", str(s)).strip().lower().rstrip(":").strip()


def cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def as_date(v) -> tuple[str, bool]:
    """ISO date and whether it parsed. Keeps the raw text when it does not."""
    if isinstance(v, (dt.datetime, dt.date)):
        return (v.date() if isinstance(v, dt.datetime) else v).isoformat(), True
    s = cell(v)
    if not s:
        return "", True
    m = re.fullmatch(r"(\d{1,2})[-./](\d{1,2})[-./](\d{2,4})", s)
    if m:
        d, mo, y = (int(x) for x in m.groups())
        y += 2000 if y < 100 else 0
        try:
            return dt.date(y, mo, d).isoformat(), True
        except ValueError:
            return s, False
    m = re.fullmatch(r"(\d{1,2})\.?\s*([a-zæøå]+)\.?\s*(\d{4})", s.lower())
    if m and m.group(2) in DA_MONTHS:
        try:
            return dt.date(int(m.group(3)), DA_MONTHS[m.group(2)], int(m.group(1))).isoformat(), True
        except ValueError:
            return s, False
    return s, False


def frame_digits(v) -> tuple[str, str, bool]:
    """A frame's digits cell, what was stripped from it, and whether it was
    REBUILT. The register often carries a stray comma or dot before the digits
    (`WCK ,000639 L`); the owner allowed stripping exactly that (DECISIONS
    2026-10-21), each one listed. A Danish Excel reads `,000639` as the number
    0.000639 and drops trailing zeros (`,000800` → 0.0008), so a number below 1
    is rebuilt as at least six digits — a guess, which `build_plan` accepts only
    when the same frame is also written as text (DECISIONS 2026-10-27)."""
    if isinstance(v, float) and 0 < v < 1:
        dec = format(Decimal(repr(v)), "f").split(".")[1]
        return dec.ljust(6, "0"), "", True
    d = cell(v).replace(" ", "")
    m = re.fullmatch(r"([,.]+)(\d.*)", d)
    return (m.group(2), m.group(1), False) if m else (d, "", False)


def joined(parts) -> str:
    return "".join(cell(p).replace(" ", "") for p in parts if cell(p)).upper()


def find_header(rows):
    """(header row index, {col: field}) — the row with most recognised labels
    that includes the frame column; labels one row above (EAN, 'aftale') merge in."""
    best = None
    for i, r in enumerate(rows[:12]):
        hits = {j: LABELS[norm(c)] for j, c in enumerate(r or ()) if c is not None and norm(c) in LABELS}
        if "frame" in hits.values() and (best is None or len(hits) > len(best[1])):
            best = (i, hits)
    if best and best[0] > 0:
        for j, c in enumerate(rows[best[0] - 1] or ()):
            if c is not None and norm(c) in ("ean", "aftale") and j not in best[1]:
                best[1][j] = LABELS[norm(c)]
    return best


# ---------------------------------------------------------------- production

def prod_query(sql: str):
    out = subprocess.run(["supabase", "db", "query", "--linked", sql], cwd=REPO,
                         capture_output=True, text=True, check=True).stdout
    data = json.loads(out[out.index("{"):])
    row = data["rows"][0]
    val = next(iter(row.values()))
    return val if isinstance(val, (list, dict)) else json.loads(val or "null")


def load_production():
    orgs = prod_query(
        "select coalesce(json_agg(x), '[]'::json) from (select o.id, o.legal_name, "
        "(select coalesce(json_agg(json_build_object('id', u.id, 'name', u.name)), '[]'::json) "
        " from organization_units u where u.organization_id = o.id and u.deleted_at is null) units "
        "from organizations o where o.deleted_at is null) x")
    frames = prod_query("select coalesce(json_agg(upper(replace(frame_number, ' ', ''))), '[]'::json) "
                        "from bikes where deleted_at is null")
    idents = prod_query(
        "select coalesce(json_agg(json_build_object('slug', t.slug, 'value', upper(replace(i.identifier_value, ' ', ''))))"
        ", '[]'::json) from bike_identifiers i join bike_identifier_types t on t.id = i.identifier_type_id "
        "where i.is_active and t.slug in ('battery_number', 'charger_number', 'fleet_number')")
    return orgs, set(frames), {(i["slug"], i["value"]) for i in idents}


# ---------------------------------------------------------------- matching

def key(s: str) -> str:
    s = s.lower().replace("tåstrup", "taastrup").replace("tårbæk", "taarbæk").replace("-", " ").replace(".", " ")
    s = re.sub(r"\b(kommune|a/s|aps|as|i/s)\b", " ", s)
    return " ".join(re.sub(r"[^a-z0-9æøå ]", " ", s).split())


def score(a: str, b: str) -> float:
    ka, kb = key(a), key(b)
    if not ka or not kb:
        return 0.0
    r = difflib.SequenceMatcher(None, ka, kb).ratio()
    wa, wb = set(ka.split()), set(kb.split())
    return max(r, len(wa & wb) / max(len(wa), len(wb)))


def top(target: str, cands, n=3):
    return sorted(((round(score(target, label), 2), cid, label) for cid, label in cands), reverse=True)[:n]


# ---------------------------------------------------------------- parsing

def parse_fleet(wb):
    bikes, skipped = [], []
    for ws in wb.worksheets:
        if ws.title in NOT_FLEET:
            continue
        rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
        found = find_header(rows)
        if not found:
            skipped.append((ws.title, "no header with a frame-number column"))
            continue
        hi, cols = found
        if "_personal" in cols.values():
            skipped.append((ws.title, "personal data (private individuals) — not imported"))
            continue
        known = set(cols)
        for k in [c for c, f in cols.items() if f in SPLIT3]:
            known |= {k + 1, k + 2}
        frame_col = next(c for c, f in cols.items() if f == "frame")
        for ri in range(hi + 1, len(rows)):
            r = rows[ri] or ()
            get = lambda j: r[j] if j < len(r) else None  # noqa: E731
            digits, stripped, rebuilt = frame_digits(get(frame_col + 1))
            if not digits:
                continue
            rec = {"sheet": ws.title.strip(), "row": ri + 1, "issues": [], "frame_stripped": stripped,
                   "frame_rebuilt": rebuilt}
            prefix, year = cell(get(frame_col)).upper(), cell(get(frame_col + 2)).upper()
            rec["frame"] = f"{prefix}{digits}{year}".upper().replace(" ", "")
            as_written = cell(get(frame_col + 1)) if rebuilt else stripped + digits
            rec["frame_display"] = " ".join(x for x in (prefix, as_written, year) if x)
            if not re.fullmatch(r"[A-Z]", year or "-"):
                rec["issues"].append("frame year letter missing or odd" if not year else f"frame suffix '{year}'")
            invoiced = []
            for j, f in cols.items():
                if f in ("frame",) or f.startswith("_"):
                    continue
                if f in SPLIT3:
                    rec[f] = joined([get(j), get(j + 1), get(j + 2)])
                elif f in ("delivered", "takeover"):
                    rec[f], ok = as_date(get(j))
                    if not ok:
                        rec["issues"].append(f"{f} date unreadable: {rec[f]}")
                elif f == "invoiced":
                    d, _ = as_date(get(j))
                    done = cell(get(j + 1)).upper() == "X"
                    if d:
                        invoiced.append(f"{d}{' X' if done else ''}")
                elif f == "text":
                    rec.setdefault("notes", [])
                    if cell(get(j)):
                        rec["notes"].append(cell(get(j)))
                else:
                    rec[f] = cell(get(j))
            # free text in columns nobody labelled (repair and battery history live there)
            extra = [cell(get(j)) for j in range(len(r)) if j not in known and j != 0 and cell(get(j))
                     and not re.fullmatch(r"X|\d{1,3}", cell(get(j)))]
            rec["notes"] = rec.get("notes", []) + [e for e in extra if len(e) > 3]
            rec["service_invoiced"] = invoiced
            a = cell(get(0))
            rec["recognition_code"] = a.upper() if RECOGNITION.match(a.upper()) else ""
            if a and not rec["recognition_code"]:
                rec["notes"].insert(0, f"col A: {a}")
            text = " ".join([a, rec.get("agreement", "")] + rec["notes"]).lower()
            rec["status"], rec["status_reason"] = "in_service", ""
            for pat, status, why in STATUS_WORDS:
                if re.search(pat, text):
                    rec["status"], rec["status_reason"] = status, why
                    break
            ag = rec.get("agreement", "").lower()
            rec["agreement"] = ("cancelled" if "opsagt" in ag else "yes" if ag.startswith("j")
                                else "no" if ag.startswith("n") else "unknown" if not ag else ag)
            bikes.append(rec)
    return bikes, skipped


def parse_schedule(wb):
    out = []
    for m in MONTHS:
        if m not in wb.sheetnames:
            continue  # a file holding customer sheets only
        ws = wb[m]
        rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
        hi = next((i for i, r in enumerate(rows[:6]) if r and any(norm(c).startswith("stelnummer") for c in r if c)), None)
        if hi is None:
            continue
        h = [norm(c) if c else "" for c in rows[hi]]
        col = lambda pred: next((i for i, x in enumerate(h) if pred(x)), None)  # noqa: E731
        fi, pi, ki = col(lambda x: x.startswith("stelnummer")), col(lambda x: x.startswith("pris")), col(lambda x: x == "kommune")
        ci, bi, ei = col(lambda x: x == "cykel nr"), col(lambda x: x.startswith("købt")), col(lambda x: x.startswith("ean"))
        years = [(i, x.split()[0]) for i, x in enumerate(h) if re.fullmatch(r"\d års", x)]
        for ri in range(hi + 1, len(rows)):
            r = rows[ri] or ()
            get = lambda j: r[j] if j is not None and j < len(r) else None  # noqa: E731
            digits, stripped, rebuilt = frame_digits(get(fi + 1))
            if not digits:
                continue
            yrs = []
            for i, n in years:
                d, _ = as_date(get(i))
                if d:
                    yrs.append(f"year {n}: {d}{' X' if cell(get(i + 1)).upper() == 'X' else ''}")
            out.append({
                "month": m, "row": ri + 1,
                "frame": f"{cell(get(fi))}{digits}{cell(get(fi + 2))}".upper().replace(" ", ""),
                "recognition_code": cell(get(ci)).upper(), "customer_text": cell(get(ki)),
                "bought": as_date(get(bi))[0], "ean": cell(get(ei)),
                "renewals": "; ".join(yrs), "yearly_price": cell(get(pi)),
                "frame_stripped": stripped, "frame_rebuilt": rebuilt,
            })
    return out


# ---------------------------------------------------------------- review

def review(src: Path, out: Path):
    import openpyxl
    print(f"reading {src.name} …", file=sys.stderr)
    wb = openpyxl.load_workbook(src, data_only=True)
    bikes, skipped = parse_fleet(wb)
    schedule = parse_schedule(wb)
    print("reading production (read-only) …", file=sys.stderr)
    orgs, prod_frames, prod_idents = load_production()
    org_cands = [(o["id"], o["legal_name"]) for o in orgs]
    by_id = {o["id"]: o for o in orgs}

    # customers: sheet -> org; (sheet, department) -> unit or org
    sheet_best = {s: top(s, org_cands, 3) for s in {b["sheet"] for b in bikes}}
    pairs = Counter((b["sheet"], b.get("department", "")) for b in bikes)
    cust_rows, pair_pick = [], {}
    for (sheet, dept), n in sorted(pairs.items()):
        sb = sheet_best[sheet]
        org = by_id.get(sb[0][1]) if sb else None
        dcands = [(("unit", u["id"]), f"{org['legal_name']} › {u['name']}") for u in (org["units"] if org else [])]
        dcands += [(("org", o["id"]), o["legal_name"]) for o in orgs]
        db = top(f"{sheet} {dept}", dcands, 3) if dept else []
        eans = sorted({b.get("ean", "") for b in bikes if b["sheet"] == sheet and b.get("department", "") == dept} - {""})
        strong_sheet = bool(sb) and sb[0][0] >= 0.9
        strong_dept = bool(db) and db[0][0] >= 0.85
        pair_pick[(sheet, dept)] = (sb[0] if sb else None, db[0] if db else None)
        cust_rows.append({
            "sheet": sheet, "department": dept, "bikes": n, "ean": " ".join(eans),
            "customer_1": f"{sb[0][2]} ({sb[0][0]})" if sb else "",
            "customer_2": f"{sb[1][2]} ({sb[1][0]})" if len(sb) > 1 else "",
            "department_1": f"{db[0][2]} ({db[0][0]})" if db else "",
            "department_2": f"{db[1][2]} ({db[1][0]})" if len(db) > 1 else "",
            "check": "ok" if strong_sheet and (not dept or strong_dept) else
                     "customer?" if not strong_sheet else "department?",
        })

    # duplicates and collisions
    by_frame = defaultdict(list)
    for b in bikes:
        by_frame[b["frame"]].append(b)
    dup_rows = []
    for f, group in by_frame.items():
        if len(group) > 1:
            for b in group:
                dup_rows.append({"frame": f, "sheet": b["sheet"], "row": b["row"],
                                 "recognition_code": b["recognition_code"], "department": b.get("department", ""),
                                 "delivered": b.get("delivered", ""), "status": b["status"]})
    ident_uses = defaultdict(list)
    for b in bikes:
        for slug, fld in (("battery_number", "battery"), ("charger_number", "charger")):
            v = b.get(fld, "")
            if re.search(r"\d{6,}", v):
                ident_uses[(slug, v)].append(b)
    conflict_rows = []
    model_codes = set()
    for (slug, v), group in ident_uses.items():
        frames = {b["frame"] for b in group}
        # Newer bikes carry the charger's MODEL code (FY2010001 on dozens of bikes)
        # where older ones carry a serial: a model code is kept in the source row
        # only, never imported as a unique charger number.
        if slug == "charger_number" and re.fullmatch(r"(AWC)?FY\d+Z?", v):
            model_codes.add(v)
            continue
        if len(frames) > 1:
            for b in group:
                conflict_rows.append({"identifier": slug, "value": v, "frame": b["frame"], "sheet": b["sheet"], "row": b["row"],
                                      "why": "same number on several bikes (the database allows it once)"})
        if (slug, v) in prod_idents:
            conflict_rows.append({"identifier": slug, "value": v, "frame": group[0]["frame"], "sheet": group[0]["sheet"],
                                  "row": group[0]["row"], "why": "already on a bike in production"})
    for b in bikes:
        if b["frame"] in prod_frames:
            conflict_rows.append({"identifier": "frame_number", "value": b["frame"], "frame": b["frame"], "sheet": b["sheet"],
                                  "row": b["row"], "why": "frame already in production"})

    # schedule <-> fleet
    sched_by_frame = defaultdict(list)
    for s in schedule:
        sched_by_frame[s["frame"]].append(s)
        s["in_fleet_sheets"] = "yes" if s["frame"] in by_frame else "no"

    out.mkdir(parents=True, exist_ok=True)
    bike_rows = []
    for b in bikes:
        pick_org, pick_dept = pair_pick[(b["sheet"], b.get("department", ""))]
        sch = sched_by_frame.get(b["frame"], [])
        bike_rows.append({
            "sheet": b["sheet"], "row": b["row"], "recognition_code": b["recognition_code"],
            "frame": b["frame"], "frame_as_written": b["frame_display"], "model": b.get("model", ""),
            "gender": b.get("gender", ""), "size": b.get("size", ""), "battery": b.get("battery", ""),
            "charger": b.get("charger", ""), "key_no": b.get("key_no", ""), "battery_key_no": b.get("battery_key_no", ""),
            "gps": b.get("gps", ""), "gps_phone": b.get("gps_phone", ""), "delivered": b.get("delivered", ""),
            "takeover": b.get("takeover", ""), "department": b.get("department", ""), "site": b.get("site", ""),
            "ean": b.get("ean", ""), "agreement_flag": b["agreement"],
            "in_renewal_schedule": "yes" if sch else "no",
            "yearly_price": " / ".join(s["yearly_price"] for s in sch),
            "renewals": " | ".join(s["renewals"] for s in sch if s["renewals"]),
            "service_invoiced": " | ".join(b["service_invoiced"]),
            "status": b["status"], "status_reason": b["status_reason"],
            "customer": pick_org[2] if pick_org else "", "department_match": pick_dept[2] if pick_dept else "",
            "notes": " · ".join(b["notes"]), "issues": " · ".join(b["issues"]),
            "listed_twice": "yes" if len(by_frame[b["frame"]]) > 1 else "",
        })

    def write(name, rows):
        with open(out / name, "w", newline="", encoding="utf-8-sig") as fh:  # BOM: Excel reads æøå
            if rows:
                w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()), delimiter=";")
                w.writeheader()
                w.writerows(rows)
        return len(rows)

    n_b = write("1-bikes.csv", bike_rows)
    n_c = write("2-customers-and-departments.csv", cust_rows)
    n_d = write("3-frames-listed-twice.csv", dup_rows)
    n_x = write("4-number-conflicts.csv", conflict_rows)
    n_s = write("5-renewal-schedule.csv", schedule)

    st = Counter(b["status"] for b in bikes)
    ag = Counter(b["agreement"] for b in bikes)
    chk = Counter(r["check"] for r in cust_rows)
    prices = Counter(s["yearly_price"] for s in schedule)
    summary = f"""# Fleet import — review ({dt.date.today().isoformat()})

Source: `{src.name}`. Nothing has been written to any database.

| | |
|---|---|
| Bike rows (customer sheets) | {n_b} — {len(by_frame)} distinct frames |
| Frames listed more than once | {len({r['frame'] for r in dup_rows})} (file 3) |
| Status from the register's words | in service {st['in_service']} · lost/stolen {st['lost_or_stolen']} · retired {st['retired']} |
| Agreement flag in the fleet sheets | yes {ag['yes']} · no {ag['no']} · unknown {ag['unknown']} · cancelled {ag['cancelled']} |
| Renewal schedule rows (monthly sheets) | {n_s} — {sum(1 for s in schedule if s['in_fleet_sheets'] == 'yes')} match a fleet-sheet bike |
| Yearly prices in the schedule | {', '.join(f'{p or "blank"} ×{c}' for p, c in prices.most_common(8))} |
| Customer / department pairs | {n_c} — ok {chk['ok']} · customer? {chk['customer?']} · department? {chk['department?']} (file 2) |
| Number conflicts | {n_x} rows (file 4) |
| Charger model codes (not serials; kept in the source row only) | {', '.join(sorted(model_codes)) or 'none'} |

## Skipped sheets
{chr(10).join(f'- **{s}** — {why}' for s, why in skipped) or '- none'}

## What to check
1. **File 2** — the customer (and department) each sheet's bikes go to. Anything not `ok` needs a
   decision: pick the right customer, or say it must be created (WOLT, Aleris and KL are not in
   the system).
2. **File 3** — frames listed twice: moved between customers, replaced, or a typo? Only one row
   per frame can be imported.
3. **File 4** — battery/charger numbers on more than one bike (the database takes each once).
4. **File 1, column `status`** — derived from words like *stjålet* / *udgået*; everything else is
   in service. Blank agreement flags are common; the renewal schedule (file 5) is what gets billed.
"""
    (out / "0-summary.md").write_text(summary, encoding="utf-8")
    print(summary)
    print(f"→ {out}", file=sys.stderr)


# ---------------------------------------------------------------- batches: plan + sql
#
# docs/plan-register-import.md §2–§7. `plan` reads the register and a database
# (read-only) and writes, for ONE batch: the mapping file a person checks
# (sheet → customer, department → unit), the exception list, the manifest of
# what the batch will create, and a review document for Dennis. `sql` re-runs
# the same planning — so nothing stale can be applied — refuses while the
# mapping is unconfirmed or anything blocks, and writes the apply / dry-run /
# verify / rollback files. Every row the batch creates gets a UUID derived from
# the batch label and its source key, so the apply file, the rollback file and
# the review name the same rows. Rows are never guessed into existence: what
# does not fit goes on the exception list, not into the database.

BATCH_NS = uuid.UUID("3b8f2c4e-9a1d-4e57-8c6b-2f0d7a1e5b93")
BATCH_DIR = REPO / "scripts" / "import-batches"
FLEET_DIR = Path.home() / "Documents/1-Projects/Jensen/Fleet"
GPS_PRICE = Decimal("2184")      # 1 704 service + 480 GPS (handling doc §7, items 3 and 6)
GPS_ONLY_PRICE = Decimal("480")  # GPS without service: decision D1, not importable yet
CHARGER_MODEL = re.compile(r"(AWC)?FY\d+Z?")
YEAR_LETTER = re.compile(r"[G-Z]")  # G = 2012 … Z = 2026
EMPTY = {"", "-", "X", "?", "NA", "N/A", "INGEN"}


def rid(label: str, kind: str, key: str) -> str:
    return str(uuid.uuid5(BATCH_NS, f"{label}|{kind}|{key}"))


def slugify(s: str) -> str:
    s = s.lower().replace("æ", "ae").replace("ø", "o").replace("å", "aa")
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


def db_json(sql: str, target: str):
    """One JSON value from the database: `prod` through the linked project
    (read-only use only), `local` through the local container's psql."""
    if target == "local":
        out = subprocess.run(["docker", "exec", "-i", "supabase_db_jensen-fms", "psql", "-U", "postgres",
                              "-d", "postgres", "-Atc", sql], capture_output=True, text=True, check=True).stdout
        return json.loads(out.strip() or "null")
    return prod_query(sql)


def load_batch_context(target: str) -> dict:
    return db_json("""select json_build_object(
      'orgs', (select coalesce(json_agg(json_build_object('id', o.id, 'name', o.legal_name, 'prefix', o.recognition_prefix,
          'units', (select coalesce(json_agg(json_build_object('id', u.id, 'name', u.name, 'ean', u.ean_number,
                    'code', u.code, 'no', u.external_customer_no) order by u.name), '[]'::json)
                    from organization_units u where u.organization_id = o.id and u.deleted_at is null))), '[]'::json)
          from organizations o where o.deleted_at is null),
      'frames', (select coalesce(json_agg(distinct f), '[]'::json) from (
          select upper(replace(frame_number, ' ', '')) f from bikes
          union select upper(replace(i.identifier_value, ' ', '')) from bike_identifiers i
            join bike_identifier_types t on t.id = i.identifier_type_id
           where t.slug = 'frame_number' and i.is_active) x),
      'idents', (select coalesce(json_agg(json_build_object('slug', t.slug, 'value', upper(replace(i.identifier_value, ' ', '')))), '[]'::json)
          from bike_identifiers i join bike_identifier_types t on t.id = i.identifier_type_id
         where i.is_active and t.slug in ('battery_number', 'charger_number')),
      'agreements', (select coalesce(json_agg(json_build_object('org', organization_id, 'unit', organization_unit_id,
          'status', status)), '[]'::json) from service_agreements where status <> 'cancelled'),
      'batches', (select coalesce(json_agg(label), '[]'::json) from import_batches),
      'people', (select coalesce(json_agg(json_build_object('id', id, 'name', full_name)), '[]'::json) from people)
    )""", target)


def parse_price(s: str):
    """Decimal, or None when blank/unreadable. '1.704', '1704,00', '2184 kr' all read."""
    t = re.sub(r"(?i)\s*(kr\.?|dkk|,-)\s*", "", s.strip())
    if not t:
        return None
    if re.fullmatch(r"\d{1,3}(\.\d{3})+(,\d+)?", t):
        t = t.replace(".", "")
    t = t.replace(",", ".")
    try:
        return Decimal(t)
    except InvalidOperation:
        return None


def iso_or_none(s: str):
    return s if re.fullmatch(r"\d{4}-\d{2}-\d{2}", s or "") else None


def present(v) -> str:
    v = cell(v).replace(" ", "").upper()
    return "" if v in EMPTY else v


class Plan:
    def __init__(self, label):
        self.label = label
        self.batch_id = rid(label, "batch", label)
        self.exceptions = []      # {severity, kind, sheet, row, frame, detail}
        self.units_new = []       # {id, org, name, ean, key}
        self.unit_fills = []      # {id, ean, name}
        self.bikes = []           # {id, frame, type, status, org, unit, assigned, site, row}
        self.idents = []          # {id, bike, slug, value, active, notes}
        self.agreements = []      # {id, org, unit, name_en, name_da, status, start, gps, row}
        self.lines = []           # {id, agreement, bike, start, price, gps, status, row}

    def flag(self, severity, kind, detail, src=None, frame=""):
        self.exceptions.append({"severity": severity, "kind": kind, "sheet": (src or {}).get("sheet", (src or {}).get("month", "")),
                                "row": (src or {}).get("row", ""), "frame": frame or (src or {}).get("frame", ""),
                                "detail": detail})

    def manifest(self) -> dict:
        c = Counter
        return {
            "batch": self.label, "batch_id": self.batch_id,
            "units_created": len(self.units_new), "unit_eans_filled": len(self.unit_fills),
            "bikes": len(self.bikes), "bikes_by_status": dict(c(b["status"] for b in self.bikes)),
            "identifiers_active": dict(c(i["slug"] for i in self.idents if i["active"])),
            "identifiers_inactive": dict(c(i["slug"] for i in self.idents if not i["active"])),
            "agreements": len(self.agreements), "agreements_by_status": dict(c(a["status"] for a in self.agreements)),
            "lines_active": sum(1 for x in self.lines if x["status"] == "active"),
            "lines_active_gps": sum(1 for x in self.lines if x["status"] == "active" and x["gps"]),
            "lines_ended": sum(1 for x in self.lines if x["status"] == "ended"),
            "exceptions": dict(c(e["severity"] for e in self.exceptions)),
        }


def merge_mapping(label, src, sheets, sel, ctx, old):
    """The mapping a person checks. Proposals are refreshed; anything already
    `confirmed` is kept exactly as the person left it."""
    orgs = {o["id"]: o for o in ctx["orgs"]}
    cands = [(o["id"], o["name"]) for o in ctx["orgs"]]
    old_c = {c["sheet"]: c for c in (old or {}).get("customers", [])}
    old_d = {(d["sheet"], d["department"], d["ean"]): d for d in (old or {}).get("departments", [])}
    customers, departments = [], []
    for sheet in sheets:
        if old_c.get(sheet, {}).get("confirmed"):
            customers.append(old_c[sheet])
            org_id = old_c[sheet]["organization_id"]
        else:
            best = top(sheet, cands, 3)
            org_id = best[0][1] if best and best[0][0] >= 0.9 else None
            customers.append({"sheet": sheet, "organization_id": org_id,
                              "organization_name": orgs[org_id]["name"] if org_id else None,
                              "proposals": [{"id": i, "name": n, "score": s} for s, i, n in best],
                              "confirmed": False})
        org = orgs.get(org_id)
        pairs = Counter((b.get("department", ""), present(b.get("ean", ""))) for b in sel if b["sheet"] == sheet)
        for (dept, ean), n in sorted(pairs.items()):
            k = (sheet, dept, ean)
            units = org["units"] if org else []
            if old_d.get(k, {}).get("confirmed"):
                d = dict(old_d[k])
                d["bikes"] = n
                departments.append(d)
                continue
            by_ean = [u for u in units if ean and u["ean"] == ean]
            scored = sorted(((round(score(dept, u["name"]), 2), u) for u in units), key=lambda x: -x[0]) if dept else []
            if by_ean:
                pick, why = by_ean[0]["id"], "same EAN"
            elif scored and scored[0][0] >= 0.85:
                pick, why = scored[0][1]["id"], f"name {scored[0][0]}"
            elif not dept and not ean:
                pick, why = "customer", "no department and no EAN"
            else:
                pick, why = None, "undecided"
            departments.append({
                "sheet": sheet, "department": dept, "ean": ean, "bikes": n,
                "unit": pick, "why": why, "new_name": dept or None,
                "candidates": [{"id": u["id"], "name": u["name"], "ean": u["ean"], "no": u["no"]} for u in units],
                "confirmed": False,
            })
    return {"batch": label, "source": src.name, "updated": dt.date.today().isoformat(),
            "how": ("Set each customer's organization_id and each department's `unit` to a unit id from "
                    "`candidates`, \"new\" (creates a department named new_name with this EAN) or "
                    "\"customer\" (the bikes go on the customer, no department), then confirmed: true. "
                    "`sql` refuses while anything is unconfirmed."),
            "customers": customers, "departments": departments}


def build_plan(args):
    import openpyxl
    label = args.batch
    if not re.fullmatch(r"[a-z0-9-]+", label):
        sys.exit("--batch must be lowercase letters, digits and dashes (e.g. register-2026-10-pilot-allerod)")
    print(f"reading {args.src.name} …", file=sys.stderr)
    wb = openpyxl.load_workbook(args.src, data_only=True)
    all_bikes, _ = parse_fleet(wb)
    schedule = parse_schedule(wb)
    by_name = {b["sheet"].upper(): b["sheet"] for b in all_bikes}
    sheets = []
    for s in args.sheets:
        if s.upper() not in by_name:
            sys.exit(f"no customer sheet called {s!r}; sheets: {', '.join(sorted(by_name.values()))}")
        sheets.append(by_name[s.upper()])
    sel = [b for b in all_bikes if b["sheet"] in sheets]
    print(f"reading the {args.db} database (read-only) …", file=sys.stderr)
    ctx = load_batch_context(args.db)
    if label in ctx["batches"]:
        sys.exit(f"batch {label!r} is already in the {args.db} database")

    map_path = BATCH_DIR / label / "mapping.json"
    old = json.loads(map_path.read_text(encoding="utf-8")) if map_path.exists() else None
    mapping = merge_mapping(label, args.src, sheets, sel, ctx, old)
    map_path.parent.mkdir(parents=True, exist_ok=True)
    map_path.write_text(json.dumps(mapping, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    P = Plan(label)
    orgs = {o["id"]: o for o in ctx["orgs"]}
    units = {u["id"]: dict(u, org=o["id"]) for o in ctx["orgs"] for u in o["units"]}
    cust = {c["sheet"]: c for c in mapping["customers"]}
    dept = {(d["sheet"], d["department"], d["ean"]): d for d in mapping["departments"]}
    for c in mapping["customers"]:
        if not c["confirmed"] or c["organization_id"] not in orgs:
            P.flag("blocks", "customer not confirmed", f"sheet {c['sheet']}: confirm organization_id in {map_path.name}")

    # ---- departments → unit ids (created, matched, or the customer itself)
    unit_for, unit_eans = {}, defaultdict(set)
    for d in mapping["departments"]:
        k = (d["sheet"], d["department"], d["ean"])
        org_id = cust[d["sheet"]]["organization_id"]
        if not d["confirmed"] or d["unit"] is None:
            P.flag("blocks", "department not confirmed",
                   f"{d['sheet']} · '{d['department'] or '(blank)'}' · EAN {d['ean'] or '(none)'} · {d['bikes']} bikes")
            continue
        if d["unit"] == "customer":
            unit_for[k] = None
        elif d["unit"] == "new":
            uid = rid(label, "unit", "|".join(k))
            P.units_new.append({"id": uid, "org": org_id, "name": d["new_name"] or d["department"] or "Afdeling",
                                "ean": d["ean"] or None, "key": "|".join(k)})
            unit_for[k] = uid
        elif d["unit"] in units and units[d["unit"]]["org"] == org_id:
            unit_for[k] = d["unit"]
            if d["ean"]:
                unit_eans[d["unit"]].add(d["ean"])
        else:
            P.flag("blocks", "department maps to a unit of another customer", f"{d['sheet']} · {d['department']} → {d['unit']}")
        if not d["ean"]:
            P.flag("note", "department without EAN", f"{d['sheet']} · '{d['department'] or '(blank)'}' · {d['bikes']} bikes")
    for uid, eans in unit_eans.items():
        u = units[uid]
        if len(eans) > 1:
            P.flag("note", "department with two EANs", f"{u['name']}: {', '.join(sorted(eans))} — none written")
        elif u["ean"] and u["ean"] not in eans:
            P.flag("note", "EAN differs from the system", f"{u['name']}: system {u['ean']}, sheet {next(iter(eans))} — not overwritten")
        elif not u["ean"]:
            P.unit_fills.append({"id": uid, "ean": next(iter(eans)), "name": u["name"]})

    # ---- frames: valid, new to the database, once in the register
    used = set(ctx["frames"])
    everywhere = defaultdict(list)
    for b in all_bikes:
        everywhere[b["frame"]].append(b)
    written = text_written_frames(all_bikes, schedule)
    chosen = {}
    for b in sel:
        f = b["frame"]
        if re.search(r"[^A-Z0-9]", f) or not YEAR_LETTER.fullmatch(f[-1:]) or b["issues"] and any("frame" in i for i in b["issues"]):
            P.flag("skipped", "frame unreadable", f"as written: {b['frame_display']}", b, f)
            continue
        if b["frame_rebuilt"]:
            ok, detail = check_rebuilt(b, written, schedule)
            if not ok:
                P.flag("skipped", "frame rebuilt from an Excel number, unconfirmed", detail, b, f)
                continue
            P.flag("note", "frame rebuilt from an Excel number", detail, b, f)
        if f in used:
            P.flag("skipped", "frame already in the system", "not imported again", b, f)
            continue
        elsewhere = sorted({x["sheet"] for x in everywhere[f]} - set(sheets))
        if elsewhere:
            P.flag("skipped", "frame also on another customer's sheet", f"also on {', '.join(elsewhere)} — moved bike?", b, f)
            continue
        if b["frame_stripped"]:
            P.flag("note", "frame punctuation removed", f"as written: {b['frame_display']} — read as {f}", b, f)
        if f in chosen:
            keep, drop = sorted([chosen[f], b], key=lambda x: x.get("delivered", "") or "", reverse=True)
            chosen[f] = keep
            keep.setdefault("also_listed", []).append({"sheet": drop["sheet"], "row": drop["row"]})
            P.flag("note", "frame listed twice", f"row {keep['row']} used (latest delivery), row {drop['row']} kept as history", drop, f)
            continue
        chosen[f] = b

    # ---- bikes and identifiers
    bike_of = {}
    battery_or_charger = defaultdict(list)
    prod_idents = {(i["slug"], i["value"]) for i in ctx["idents"]}
    for f, b in chosen.items():
        k = (b["sheet"], b.get("department", ""), present(b.get("ean", "")))
        if k not in unit_for:
            continue  # its department is unconfirmed — already flagged
        org_id = cust[b["sheet"]]["organization_id"]
        bid = rid(label, "bike", f)
        row = {x: v for x, v in b.items() if x not in ("issues",)}
        P.bikes.append({"id": bid, "frame": f, "type": "e_bike" if present(b.get("battery")) else "pedal",
                        "status": b["status"], "org": org_id, "unit": unit_for[k],
                        "assigned": iso_or_none(b.get("delivered", "")), "site": b.get("site") or None, "row": row,
                        "delivered": b.get("delivered", ""), "src": b})
        bike_of[f] = P.bikes[-1]
        P.idents.append({"id": rid(label, "ident", f"{f}|frame"), "bike": bid, "slug": "frame_number",
                         "value": f, "active": True, "notes": None})
        if b["recognition_code"]:
            P.idents.append({"id": rid(label, "ident", f"{f}|code"), "bike": bid, "slug": "fleet_number",
                             "value": b["recognition_code"], "active": True, "notes": None})
        for fld, slug in (("key_no", "lock_number"), ("battery_key_no", "battery_lock_number")):
            v = present(b.get(fld))
            if v:
                P.idents.append({"id": rid(label, "ident", f"{f}|{slug}"), "bike": bid, "slug": slug,
                                 "value": v, "active": True, "notes": None})
        for fld, slug in (("battery", "battery_number"), ("charger", "charger_number")):
            v = present(b.get(fld))
            if not v or (slug == "charger_number" and CHARGER_MODEL.fullmatch(v)):
                continue  # a model code stays in import_row only
            battery_or_charger[(slug, v)].append(P.bikes[-1])
    for (slug, v), group in battery_or_charger.items():
        group = sorted(group, key=lambda x: x["delivered"] or "", reverse=True)  # newest bike keeps it (rule B3)
        for i, bk in enumerate(group):
            active, note = True, None
            if (slug, v) in prod_idents:
                active, note = False, "Same number is active on another bike in the system (import)."
                P.flag("note", f"{slug.split('_')[0]} number already in the system", f"{v} — imported inactive", bk["src"], bk["frame"])
            elif i > 0:
                active, note = False, f"Same number on a newer bike, {group[0]['frame']} (import)."
                P.flag("note", f"{slug.split('_')[0]} number on two bikes", f"{v} — active on {group[0]['frame']}", bk["src"], bk["frame"])
            P.idents.append({"id": rid(label, "ident", f"{bk['frame']}|{slug}"), "bike": bk["id"], "slug": slug,
                             "value": v, "active": active, "notes": note})

    # recognition codes: not unique by rule, a duplicate within a customer is a typo in practice
    codes = defaultdict(list)
    for bk in P.bikes:
        c = bk["src"]["recognition_code"]
        if c:
            codes[(bk["org"], c)].append(bk)
        prefix = orgs[bk["org"]]["prefix"]
        if c and prefix and not c.startswith(prefix):
            P.flag("note", "code does not start with the customer's prefix", f"{c} (prefix {prefix})", bk["src"], bk["frame"])
    for (_, c), group in codes.items():
        if len(group) > 1:
            P.flag("note", "same code on two bikes", f"{c}: {', '.join(x['frame'] for x in group)}", group[0]["src"])

    # ---- the renewal schedule → agreement lines
    sel_frames = {b["frame"] for b in sel}
    fleet_frames = set(everywhere)
    sched = defaultdict(list)
    def for_these_sheets(text):
        words = set(key(text).split())
        return any(score(text, sh) >= 0.9 or (set(key(sh).split()) and set(key(sh).split()) <= words) for sh in sheets)

    def one_off(a, b):  # same length, one character apart — a typo, most likely
        return len(a) == len(b) and sum(x != y for x, y in zip(a, b)) == 1

    for s in schedule:
        if s["frame"] in sel_frames:
            sched[s["frame"]].append(s)
            if s["frame_stripped"]:
                P.flag("note", "frame punctuation removed", f"renewal sheet: '{s['frame_stripped']}' before the digits — read as {s['frame']}", s)
            if s["frame_rebuilt"]:
                P.flag("note", "frame rebuilt from an Excel number", f"renewal sheet: read as {s['frame']}", s)
        elif s["frame"] not in fleet_frames and for_these_sheets(s["customer_text"]):
            near = [f for f in sorted(sel_frames) if one_off(f, s["frame"])]
            P.flag("skipped", "renewal row with no bike", f"{s['customer_text']} · price {s['yearly_price'] or 'blank'}"
                   + (" · frame rebuilt from an Excel number" if s["frame_rebuilt"] else "")
                   + (f" — nearest frame on the sheet: {', '.join(near)} (a typo?)" if near else
                      " — add the bike to the customer's sheet, or confirm it is gone"), s)
    today = dt.date.today().isoformat()
    lines_by_key = defaultdict(list)
    for bk in P.bikes:
        rows = sched.get(bk["frame"], [])
        b = bk["src"]
        if not rows:
            P.flag("note", "bike with no renewal row", ("marked 'ja' on the sheet, but " if b["agreement"] == "yes" else "")
                   + "no line written — imported as not covered", b, bk["frame"])
            continue
        priced, zero = [], []
        for s in rows:
            p = parse_price(s["yearly_price"])
            if p is None:
                P.flag("skipped", "renewal row without a price", f"price '{s['yearly_price']}' — no line", s)
            elif p == 0:
                zero.append(s)
            elif p == GPS_ONLY_PRICE:
                P.flag("skipped", "GPS-only row (decision D1)", "480 kr — not importable until GPS-only is decided", s)
            else:
                priced.append((s, p))
        if len(priced) > 1:
            P.flag("skipped", "two priced renewal rows", " / ".join(f"{s['month']} {p}" for s, p in priced) + " — no line", b, bk["frame"])
            priced = []
        if zero and b["agreement"] == "yes":
            P.flag("note", "0 kr row on a bike marked 'ja'", "imported as an ended line", zero[0], bk["frame"])
        if zero and priced:
            P.flag("note", "0 kr and priced rows", "an ended AND an active line — a replacement or a move?", b, bk["frame"])
        k = (bk["org"], bk["unit"])
        if zero:
            s = zero[0]
            lines_by_key[k].append({"bike": bk, "status": "ended", "price": None, "gps": False, "src": s,
                                    "start": pick_start(P, s, b, args.start_date, bk["frame"], today, quiet=True)})
        if priced:
            s, p = priced[0]
            if b["status"] != "in_service":
                P.flag("skipped", f"{b['status'].replace('_', ' ')} bike with a priced row", f"{p} kr — no line; is it still invoiced?", b, bk["frame"])
                continue
            start = pick_start(P, s, b, args.start_date, bk["frame"], today)
            if start:
                lines_by_key[k].append({"bike": bk, "status": "active", "price": p, "gps": p == GPS_PRICE, "src": s, "start": start})

    # ---- agreements: one per customer × department
    existing = {(a["org"], a["unit"]) for a in ctx["agreements"]}
    for (org_id, unit_id), group in sorted(lines_by_key.items(), key=lambda kv: (kv[0][0], kv[0][1] or "")):
        org_name = orgs[org_id]["name"]
        unit_name = (units.get(unit_id) or next((u for u in P.units_new if u["id"] == unit_id), {})).get("name") if unit_id else None
        if (org_id, unit_id) in existing:
            P.flag("blocks", "customer/department already has an agreement", f"{org_name}{', ' + unit_name if unit_name else ''}")
            continue
        group = [x for x in group if x["start"]]
        if not group:
            continue
        aid = rid(label, "agreement", f"{org_id}|{unit_id or ''}")
        active = [x for x in group if x["status"] == "active"]
        who = org_name + (f", {unit_name}" if unit_name else "")
        P.agreements.append({
            "id": aid, "org": org_id, "unit": unit_id,
            "name_en": f"Service agreement — {who}", "name_da": f"Serviceaftale — {who}",
            "status": "active" if active else "expired", "start": min(x["start"] for x in group),
            "gps": any(x["gps"] for x in active),
            "row": {"sheet": group[0]["bike"]["src"]["sheet"], "department": unit_name,
                    "eans": sorted({x["src"].get("ean", "") for x in group} - {""})},
        })
        for x in group:
            P.lines.append({"id": rid(label, "line", f"{x['bike']['frame']}|{x['status']}"), "agreement": aid,
                            "bike": x["bike"]["id"], "start": x["start"], "price": x["price"], "gps": x["gps"],
                            "status": x["status"], "row": x["src"]})
    return P, mapping, map_path, sheets


def text_written_frames(bikes, schedule) -> dict:
    """Frame → where it is written as TEXT (a customer sheet or a renewal row).
    A frame rebuilt from an Excel number counts only if it is here."""
    out = {}
    for b in bikes:
        if not b["frame_rebuilt"]:
            out.setdefault(b["frame"], f"{b['sheet']} row {b['row']}")
    for s in schedule:
        if not s["frame_rebuilt"]:
            out.setdefault(s["frame"], f"renewal sheet {s['month']} row {s['row']}")
    return out


def check_rebuilt(b, written, schedule) -> tuple[bool, str]:
    """(accepted, detail) for a customer-sheet bike whose frame was rebuilt."""
    f, held = b["frame"], b["frame_display"]
    if f in written:
        return True, f"Excel held {held} — read as {f}, confirmed by {written[f]}"
    code = b["recognition_code"]
    other = sorted({s["frame"] for s in schedule if code and s["recognition_code"] == code})
    return False, (f"Excel held {held} — read as {f}, but no row writes that frame as text"
                   + (f"; the renewal sheet has {code} as {', '.join(other)}" if other else "")
                   + " — Dennis types the digits as text")


def pick_start(P, s, b, prefer, frame, today, quiet=False):
    """A line's start: the schedule's purchase date or the fleet sheet's delivered
    date, whichever `--start-date` prefers, else the other. Disagreements are listed."""
    sched_d, fleet_d = iso_or_none(s.get("bought", "")), iso_or_none(b.get("delivered", ""))
    first, second = (sched_d, fleet_d) if prefer == "schedule" else (fleet_d, sched_d)
    start = first or second
    if not quiet and sched_d and fleet_d and sched_d != fleet_d:
        P.flag("note", "dates disagree", f"renewal sheet {sched_d}, customer sheet {fleet_d} — used {start}", b, frame)
    if not start:
        P.flag("skipped", "line without a start date", "no purchase or delivery date — no line", s, frame)
        return None
    if start > today:
        P.flag("skipped", "line starts in the future", f"{start} — no line", s, frame)
        return None
    return start


# ---------------------------------------------------------------- SQL text

def lit(v) -> str:
    if v is None:
        return "null"
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, Decimal)):
        return str(v)
    return "'" + str(v).replace("'", "''") + "'"


def jlit(v) -> str:
    return lit(json.dumps(v, ensure_ascii=False, default=str)) + "::jsonb"


def values(rows) -> str:
    return ",\n    ".join("(" + ", ".join(r) + ")" for r in rows)


def apply_body(P: Plan, m: dict, src: str, by: str | None) -> str:
    """The batch as ONE plpgsql body: guards, inserts, then count checks that
    must equal the manifest. The dry run is this body plus a closing exception."""
    B = lit(P.batch_id)
    frames = ", ".join(lit(b["frame"]) for b in P.bikes) or "null"
    keys = ", ".join(f"({lit(a['org'])}::uuid, {lit(a['unit'])}::uuid)" for a in P.agreements) or "(null::uuid, null::uuid)"
    matched_units = sorted({b["unit"] for b in P.bikes if b["unit"]} - {u["id"] for u in P.units_new})
    out = [f"""declare n int;
begin
  -- guards: refuse against a database that is not the one this batch was planned on
  if exists (select 1 from import_batches where label = {lit(P.label)}) then
    raise exception 'batch {P.label} is already applied';
  end if;
  if exists (select 1 from bikes where upper(replace(frame_number, ' ', '')) in ({frames}))
     or exists (select 1 from bike_identifiers i join bike_identifier_types t on t.id = i.identifier_type_id
                where t.slug = 'frame_number' and i.is_active and upper(replace(i.identifier_value, ' ', '')) in ({frames})) then
    raise exception 'a frame of this batch is already in the database — re-run plan';
  end if;
  if exists (select 1 from service_agreements where status <> 'cancelled'
             and (organization_id, coalesce(organization_unit_id, '00000000-0000-0000-0000-000000000000'))
                 in (select o, coalesce(u, '00000000-0000-0000-0000-000000000000') from (values {keys}) k(o, u))) then
    raise exception 'a customer/department of this batch already has an agreement — re-run plan';
  end if;
  select count(*) into n from organization_units where deleted_at is null and id in ({', '.join(lit(u) for u in matched_units) or 'null'});
  if n <> {len(matched_units)} then raise exception 'matched departments: % of {len(matched_units)} found', n; end if;

  insert into import_batches (id, label, source, notes, imported_by)
  values ({B}, {lit(P.label)}, {lit(src)}, {lit(f"Register import: {m['bikes']} bikes, {m['agreements']} agreements, {m['lines_active'] + m['lines_ended']} lines.")}, {lit(by)});
"""]
    if P.units_new:
        out.append(f"""
  insert into organization_units (id, organization_id, name, ean_number, notes, import_batch_id) values
    {values([[lit(u['id']), lit(u['org']), lit(u['name']), lit(u['ean']), lit('Created by the register import.'), B] for u in P.units_new])};
""")
    for u in P.unit_fills:
        out.append(f"  update organization_units set ean_number = {lit(u['ean'])} where id = {lit(u['id'])} and ean_number is null;\n")
    if P.bikes:
        out.append(f"""
  insert into bikes (id, bike_type_id, frame_number, status, owner_organization_id, owner_unit_id,
                     assigned_at, current_location_text, frame_number_confirmed, import_batch_id, import_row)
  select v.id::uuid, t.id, v.frame, v.status::bike_status, v.org::uuid, v.unit::uuid, v.assigned::timestamptz,
         v.site, true, {B}, v.row::jsonb
    from (values
    {values([[lit(b['id']), lit(b['type']), lit(b['frame']), lit(b['status']), lit(b['org']), lit(b['unit']),
              lit(b['assigned']), lit(b['site']), lit(json.dumps(b['row'], ensure_ascii=False, default=str))] for b in P.bikes])}
    ) v(id, type, frame, status, org, unit, assigned, site, row)
    join bike_types t on t.slug = v.type;
  get diagnostics n = row_count;
  if n <> {len(P.bikes)} then raise exception 'bikes inserted: % of {len(P.bikes)} (a bike type slug is missing?)', n; end if;

  insert into bike_identifiers (id, bike_id, identifier_type_id, identifier_value, is_active, notes, deactivated_at)
  select v.id::uuid, v.bike::uuid, t.id, v.value, v.active, v.notes, case when v.active then null else now() end
    from (values
    {values([[lit(i['id']), lit(i['bike']), lit(i['slug']), lit(i['value']), lit(i['active']), lit(i['notes'])] for i in P.idents])}
    ) v(id, bike, slug, value, active, notes)
    join bike_identifier_types t on t.slug = v.slug;
  get diagnostics n = row_count;
  if n <> {len(P.idents)} then raise exception 'identifiers inserted: % of {len(P.idents)}', n; end if;
""")
    if P.agreements:
        out.append(f"""
  insert into service_agreements (id, organization_id, organization_unit_id, name_en, name_da, status, start_date,
                                  covers_parts, covers_labor, monthly_fee, has_gps, import_batch_id, import_row) values
    {values([[lit(a['id']), lit(a['org']), lit(a['unit']), lit(a['name_en']), lit(a['name_da']),
              lit(a['status']) + '::service_agreement_status', lit(a['start']), 'true', 'true', 'null', lit(a['gps']), B,
              jlit(a['row'])] for a in P.agreements])};

  insert into service_agreement_bikes (id, agreement_id, bike_id, start_date, yearly_price, currency, has_gps, status,
                                       end_reason, source, import_batch_id, import_row) values
    {values([[lit(x['id']), lit(x['agreement']), lit(x['bike']), lit(x['start']), lit(x['price']),
              lit('DKK' if x['price'] is not None else None), lit(x['gps']), lit(x['status']),
              lit('ended' if x['status'] == 'ended' else None), "'import'", B, jlit(x['row'])] for x in P.lines])};
""")
    out.append(f"""
  -- the batch must equal its manifest
  select count(*) into n from bikes where import_batch_id = {B};
  if n <> {m['bikes']} then raise exception 'bikes: % (manifest {m['bikes']})', n; end if;
  select count(*) into n from service_agreements where import_batch_id = {B};
  if n <> {m['agreements']} then raise exception 'agreements: % (manifest {m['agreements']})', n; end if;
  select count(*) into n from service_agreement_bikes where import_batch_id = {B} and status = 'active';
  if n <> {m['lines_active']} then raise exception 'active lines: % (manifest {m['lines_active']})', n; end if;
  select count(*) into n from service_agreement_bikes where import_batch_id = {B} and status = 'ended';
  if n <> {m['lines_ended']} then raise exception 'ended lines: % (manifest {m['lines_ended']})', n; end if;
  select count(*) into n from organization_units where import_batch_id = {B};
  if n <> {m['units_created']} then raise exception 'departments created: % (manifest {m['units_created']})', n; end if;
""")
    return "".join(out)


def verify_sql(P: Plan, m: dict) -> str:
    B = lit(P.batch_id)
    want = lambda k, n, q: f"  select {lit(k)} as check_name, {n} as expected, ({q})::int as actual"  # noqa: E731
    bs = f"(select id from bikes where import_batch_id = {B})"
    checks = [
        want("batch row", 1, f"select count(*) from import_batches where id = {B}"),
        want("bikes", m["bikes"], f"select count(*) from bikes where import_batch_id = {B}"),
        *[want(f"bikes {s}", n, f"select count(*) from bikes where import_batch_id = {B} and status = '{s}'")
          for s, n in sorted(m["bikes_by_status"].items())],
        *[want(f"identifiers {s} active", n, f"select count(*) from bike_identifiers i join bike_identifier_types t on t.id = i.identifier_type_id where i.bike_id in {bs} and t.slug = '{s}' and i.is_active")
          for s, n in sorted(m["identifiers_active"].items())],
        *[want(f"identifiers {s} inactive", n, f"select count(*) from bike_identifiers i join bike_identifier_types t on t.id = i.identifier_type_id where i.bike_id in {bs} and t.slug = '{s}' and not i.is_active")
          for s, n in sorted(m["identifiers_inactive"].items())],
        want("departments created", m["units_created"], f"select count(*) from organization_units where import_batch_id = {B}"),
        want("department EANs filled", m["unit_eans_filled"],
             "select count(*) from organization_units where " +
             (" or ".join(f"(id = {lit(u['id'])} and ean_number = {lit(u['ean'])})" for u in P.unit_fills) or "false")),
        want("agreements", m["agreements"], f"select count(*) from service_agreements where import_batch_id = {B}"),
        want("lines active", m["lines_active"], f"select count(*) from service_agreement_bikes where import_batch_id = {B} and status = 'active'"),
        want("lines active with GPS", m["lines_active_gps"], f"select count(*) from service_agreement_bikes where import_batch_id = {B} and status = 'active' and has_gps"),
        want("lines ended", m["lines_ended"], f"select count(*) from service_agreement_bikes where import_batch_id = {B} and status = 'ended'"),
        # offenders — each must be zero
        want("bikes without exactly one active frame identifier = frame_number", 0,
             f"select count(*) from bikes b where b.import_batch_id = {B} and (not b.frame_number_confirmed or 1 <> (select count(*) from bike_identifiers i join bike_identifier_types t on t.id = i.identifier_type_id where i.bike_id = b.id and t.slug = 'frame_number' and i.is_active and i.identifier_value = b.frame_number))"),
        want("in-service bikes without an owner", 0, f"select count(*) from bikes where import_batch_id = {B} and status = 'in_service' and owner_organization_id is null"),
        want("charger identifiers that are model codes", 0,
             f"select count(*) from bike_identifiers i join bike_identifier_types t on t.id = i.identifier_type_id where i.bike_id in {bs} and t.slug = 'charger_number' and i.identifier_value ~ '^(AWC)?FY[0-9]+Z?$'"),
        want("active lines with no price, a future start, or a bike owned elsewhere", 0,
             f"select count(*) from service_agreement_bikes l join service_agreements a on a.id = l.agreement_id join bikes b on b.id = l.bike_id where l.import_batch_id = {B} and l.status = 'active' and (l.yearly_price is null or l.currency is null or l.start_date > current_date or b.owner_organization_id <> a.organization_id or (a.organization_unit_id is not null and b.owner_unit_id is distinct from a.organization_unit_id))"),
        want("imported agreements with a monthly fee", 0, f"select count(*) from service_agreements where import_batch_id = {B} and monthly_fee is not null"),
    ]
    return ("-- Verify " + P.label + ". Every row must read expected = actual.\n"
            "select check_name, expected, actual, case when expected = actual then 'ok' else 'WRONG' end as result from (\n"
            + "\n  union all\n".join(checks) + "\n) c;\n")


def rollback_body(P: Plan, m: dict) -> str:
    B = lit(P.batch_id)
    bs = f"(select id from bikes where import_batch_id = {B})"
    ls = f"(select id from service_agreement_bikes where import_batch_id = {B})"
    ags = f"(select id from service_agreements where import_batch_id = {B})"
    us = f"(select id from organization_units where import_batch_id = {B})"
    guards = [
        (f"select count(*) from work_orders where bike_id in {bs} or covered_by_service_agreement_bike_id in {ls} or covered_by_service_agreement_id in {ags}", "a work order"),
        (f"select count(*) from maintenance_tickets where bike_id in {bs}", "a ticket"),
        (f"select count(*) from service_order_bikes where bike_id in {bs}", "a paint order"),
        (f"select count(*) from invoice_lines where service_agreement_id in {ags}", "an invoice line"),
        (f"select count(*) from inbound_messages where matched_bike_id in {bs}", "a call or note"),
        (f"select count(*) from service_agreement_documents where agreement_id in {ags}", "an uploaded agreement paper"),
        (f"select count(*) from service_agreement_bikes where (bike_id in {bs} or agreement_id in {ags}) and import_batch_id is distinct from {B}", "an agreement line added since"),
        (f"select count(*) from service_agreement_bikes where import_batch_id = {B} and updated_at > (select imported_at from import_batches where id = {B}) + interval '1 minute'", "a line ended or changed since"),
        (f"select count(*) from offers where organization_unit_id in {us}", "an offer on a created department"),
        (f"select count(*) from sales_orders where organization_unit_id in {us}", "a sales order on a created department"),
        (f"select count(*) from service_agreements where organization_unit_id in {us} and import_batch_id is distinct from {B}", "an agreement on a created department"),
        (f"select count(*) from bikes where owner_unit_id in {us} and import_batch_id is distinct from {B}", "a bike on a created department"),
    ]
    g = "".join(f"  if ({q}) > 0 then raise exception 'refused: {why} hangs on this batch — fix forward, do not roll back'; end if;\n"
                for q, why in guards)
    return f"""declare n int;
begin
  if not exists (select 1 from import_batches where id = {B}) then raise exception 'batch {P.label} is not in this database'; end if;
{g}
  select count(*) into n from bikes where import_batch_id = {B};
  if n <> {m['bikes']} then raise exception 'bikes: % (manifest {m['bikes']}) — not the batch that was applied', n; end if;
  select count(*) into n from service_agreement_bikes where import_batch_id = {B};
  if n <> {m['lines_active'] + m['lines_ended']} then raise exception 'lines: % (manifest {m['lines_active'] + m['lines_ended']})', n; end if;

  delete from service_agreement_bikes where import_batch_id = {B};
  delete from service_agreements where import_batch_id = {B};
  delete from bikes where import_batch_id = {B};  -- identifiers, state log and parts cascade
  delete from organization_units where import_batch_id = {B};
{''.join(f"  update organization_units set ean_number = null where id = {lit(u['id'])} and ean_number = {lit(u['ean'])};{chr(10)}" for u in P.unit_fills)}  delete from organizations where import_batch_id = {B};
  delete from import_batches where id = {B};
"""


def review_md(P: Plan, m: dict, mapping: dict, ctx_orgs: dict, src_name: str) -> str:
    sev = {"blocks": "Must be decided before the import", "skipped": "Not imported — please fix in the spreadsheet",
           "note": "Imported — please check"}
    unit_name = {u["id"]: u["name"] for o in ctx_orgs.values() for u in o["units"]} | {u["id"]: u["name"] + " (new)" for u in P.units_new}
    parts = [f"# Register import — review of {', '.join(c['sheet'] for c in mapping['customers'])}\n",
             f"**{dt.date.today().strftime('%-d %B %Y')}.** From *{src_name}*. **Nothing has been written to the "
             "system yet.** Please correct the **spreadsheet**, not this paper — the import is re-run from the "
             "spreadsheet until the list at the end holds only things you accept.\n",
             "## What goes in\n", "| | |", "|---|---|",
             f"| Bikes | {m['bikes']} — " + ", ".join(f"{s.replace('_', ' ')} {n}" for s, n in sorted(m["bikes_by_status"].items())) + " |",
             f"| Departments | {m['units_created']} new, {m['unit_eans_filled']} get their EAN |",
             f"| Service agreements | {m['agreements']} (one per department) |",
             f"| Bikes under agreement | {m['lines_active']} active, of which {m['lines_active_gps']} with GPS (2 184 kr) · {m['lines_ended']} ended (0 kr) |",
             "", "## Departments\n", "| Department as written | EAN | Bikes | Goes to |", "|---|---|---|---|"]
    for d in mapping["departments"]:
        to = {"customer": "the customer itself", "new": f"new department *{d.get('new_name') or d['department']}*", None: "**not decided**"}.get(
            d["unit"], unit_name.get(d["unit"], d["unit"]))
        parts.append(f"| {d['department'] or '(blank)'} | {d['ean'] or '—'} | {d['bikes']} | {to} |")
    prices = Counter(str(x["price"]) for x in P.lines if x["status"] == "active")
    parts += ["", "## Yearly prices on the active lines\n",
              ", ".join(f"{p} kr × {n}" for p, n in prices.most_common()) or "none", ""]
    for s in ("blocks", "skipped", "note"):
        rows = [e for e in P.exceptions if e["severity"] == s]
        if not rows:
            continue
        parts += [f"## {sev[s]} ({len(rows)})\n", "| What | Sheet · row | Frame | Detail |", "|---|---|---|---|"]
        for e in sorted(rows, key=lambda e: (e["kind"], str(e["sheet"]), str(e["row"]))):
            parts.append(f"| {e['kind']} | {e['sheet']} {('· ' + str(e['row'])) if e['row'] else ''} | {e['frame']} | {e['detail']} |")
        parts.append("")
    return "\n".join(parts) + "\n"


def run_plan(args, write_sql: bool):
    P, mapping, map_path, sheets = build_plan(args)
    m = P.manifest()
    out = args.out or FLEET_DIR / f"import-{P.label}"
    out.mkdir(parents=True, exist_ok=True)
    with open(out / "exceptions.csv", "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=["severity", "kind", "sheet", "row", "frame", "detail"], delimiter=";")
        w.writeheader()
        w.writerows(P.exceptions)
    (out / "manifest.json").write_text(json.dumps(m, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    ctx_orgs = {o["id"]: o for o in load_batch_context(args.db)["orgs"]}
    (out / "review.md").write_text(review_md(P, m, mapping, ctx_orgs, args.src.name), encoding="utf-8")
    print(json.dumps(m, ensure_ascii=False, indent=2))
    print(f"mapping → {map_path}\nreview, exceptions, manifest → {out}", file=sys.stderr)
    if not write_sql:
        return
    blocking = [e for e in P.exceptions if e["severity"] == "blocks"]
    if blocking:
        sys.exit("refused — decide these first:\n  " + "\n  ".join(f"{e['kind']}: {e['detail']}" for e in blocking))
    by = None
    if args.by:
        hits = [p for p in load_batch_context(args.db)["people"] if args.by.lower() in (p["name"] or "").lower()]
        if len(hits) != 1:
            sys.exit(f"--by {args.by!r} matches {len(hits)} people")
        by = hits[0]["id"]
    body = apply_body(P, m, args.src.name, by)
    tag = "$import$"
    assert tag not in body
    (out / "apply-dryrun.sql").write_text(
        f"-- DRY RUN of {P.label}: everything, then an exception that undoes it and reports the counts.\n"
        f"do {tag}\n{body}\n  raise exception 'DRY RUN OK — {P.label}: {m['bikes']} bikes, {m['agreements']} agreements, "
        f"{m['lines_active']} active + {m['lines_ended']} ended lines, {m['units_created']} departments created';\nend\n{tag};\n",
        encoding="utf-8")
    apply_sql = f"do {tag}\n{body}\nend\n{tag};\n"
    (out / "apply.sql").write_text(f"-- {P.label} — the register import batch (docs/plan-register-import.md)\nbegin;\n{apply_sql}commit;\n", encoding="utf-8")
    rb = rollback_body(P, m)
    assert tag not in rb
    (out / "rollback-dryrun.sql").write_text(
        f"-- DRY RUN of the rollback of {P.label}.\ndo {tag}\n{rb}\n  raise exception 'ROLLBACK DRY RUN OK — {P.label}: {m['bikes']} bikes would go';\nend\n{tag};\n",
        encoding="utf-8")
    (out / "rollback.sql").write_text(f"-- Rollback of {P.label}. NOT a migration; if it is ever run, copy it into migrations/ at the next number.\n"
                                      f"begin;\ndo {tag}\n{rb}\nend\n{tag};\ncommit;\n", encoding="utf-8")
    (out / "verify.sql").write_text(verify_sql(P, m), encoding="utf-8")
    if args.migration:
        n = args.migration
        name = f"{n}_import_{P.label.replace('-', '_')}"
        mig = REPO / "migrations" / f"{name}.sql"
        if mig.exists() or any(REPO.glob(f"migrations/{n}_*.sql")):
            sys.exit(f"migration number {n} is taken")
        mig.write_text(f"-- {n} — register import batch {P.label} (docs/plan-register-import.md).\n"
                       f"-- Generated by scripts/import_fleet.py sql from {args.src.name}; mapping in "
                       f"scripts/import-batches/{P.label}/. Rollback: scripts/import-batches/{P.label}/rollback.sql.\n\n"
                       f"begin;\n{apply_sql}\ninsert into public.schema_migrations (version, name)\n"
                       f"  values ({n}, '{name}') on conflict (version) do nothing;\n\ncommit;\n", encoding="utf-8")
        for f in ("verify.sql", "rollback.sql", "rollback-dryrun.sql", "manifest.json"):
            (map_path.parent / f).write_text((out / f).read_text(encoding="utf-8"), encoding="utf-8")
        print(f"migration → {mig}", file=sys.stderr)
    print(f"apply / dry-run / verify / rollback → {out}", file=sys.stderr)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    r = sub.add_parser("review", help="parse the register and write review CSVs (no database writes)")
    r.add_argument("--src", type=Path, default=DEFAULT_SRC)
    r.add_argument("--out", type=Path, default=DEFAULT_OUT)
    for name, help_ in (("plan", "plan ONE batch: mapping file, exceptions, manifest, review (no database writes)"),
                        ("sql", "plan again, then write the apply / dry-run / verify / rollback SQL (no database writes)")):
        p = sub.add_parser(name, help=help_)
        p.add_argument("--batch", required=True, help="label, e.g. register-2026-10-pilot-allerod")
        p.add_argument("--sheets", nargs="+", required=True, help="customer sheet names in the register")
        p.add_argument("--src", type=Path, default=DEFAULT_SRC)
        p.add_argument("--db", choices=("prod", "local"), default="prod", help="database to plan against (read-only)")
        p.add_argument("--start-date", choices=("schedule", "fleet"), default="schedule",
                       help="which date starts a line when both exist: the renewal sheet's purchase date or the customer sheet's delivery date")
        p.add_argument("--out", type=Path, default=None)
        if name == "sql":
            p.add_argument("--by", help="who runs the import (a name in people)")
            p.add_argument("--migration", type=int, help="also write migrations/<N>_import_<batch>.sql and copy verify/rollback into scripts/import-batches/")
    args = ap.parse_args()
    if args.cmd == "review":
        review(args.src, args.out)
    else:
        run_plan(args, write_sql=args.cmd == "sql")


if __name__ == "__main__":
    main()
